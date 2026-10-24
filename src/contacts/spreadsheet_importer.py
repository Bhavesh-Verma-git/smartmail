"""
src/contacts/spreadsheet_importer.py
------------------------------------
Universal contact spreadsheet reader supporting CSV, modern XLSX, and legacy XLS.
Preserves source row numbering, cell strings, and order.
Provides worksheet enumeration, header detection, column mapping, and detailed diagnostics.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from src.contacts.column_mapper import (
    STANDARD_FIELDS,
    ColumnMappingResult,
    auto_map_columns,
    validate_user_mapping,
)
from src.contacts.csv_importer import ContactRecord, ImportRowError
from src.contacts.email_parser import _FREE_PROVIDERS, EmailParseResult, parse_email


class SpreadsheetError(Exception):
    """Raised for file or parsing errors in spreadsheets."""


@dataclass
class DuplicateEmailRecord:
    """Tracks duplicate email occurrences and their source row numbers."""
    email: str
    first_row: int
    duplicate_rows: list[int] = field(default_factory=list)


@dataclass
class RowValidationResult:
    """Detailed validation state for a single spreadsheet row."""
    row_number: int  # 1-based source row number in the file
    raw_values: dict[str, str]
    email: str
    contact_name: str
    company_name: str
    website: str
    role_or_job_url: str
    source_url: str
    notes: str
    # Validation flags
    is_valid: bool
    status: str  # VALID, SYNTAX_ERROR, MISSING_EMAIL, DUPLICATE
    reason: str
    is_free_provider: bool = False
    is_missing_company: bool = False
    is_missing_name: bool = False
    contact_record: ContactRecord | None = None


@dataclass
class SpreadsheetValidationSummary:
    """
    Summary counts for pre-flight validation.
    Note on Metric Categories:
    - Mutually Exclusive Row Outcomes:
        * valid_email_count + invalid_syntax_count + missing_email_count + duplicate_email_count == total_data_rows
    - Overlapping Informational Categories:
        * missing_company_count (rows lacking company information)
        * missing_name_count (rows lacking recipient name)
        * free_provider_count (rows using free email domains like gmail.com)
    """
    total_source_rows: int = 0
    total_data_rows: int = 0
    valid_email_count: int = 0
    invalid_syntax_count: int = 0
    missing_email_count: int = 0
    duplicate_email_count: int = 0
    # Overlapping / Informational
    missing_company_count: int = 0
    missing_name_count: int = 0
    free_provider_count: int = 0
    duplicate_details: list[DuplicateEmailRecord] = field(default_factory=list)


@dataclass
class SpreadsheetInspection:
    """Initial inspection result of a spreadsheet file before processing."""
    file_type: str  # 'csv', 'xlsx', 'xls'
    sheet_names: list[str]
    selected_sheet: str
    headers: list[str]
    header_row_index: int
    total_rows_approx: int
    proposed_mapping: ColumnMappingResult


@dataclass
class SpreadsheetParseResult:
    """Complete result of parsing a spreadsheet with applied column mappings."""
    inspection: SpreadsheetInspection
    applied_mapping: dict[str, str]
    summary: SpreadsheetValidationSummary
    rows: list[RowValidationResult]
    valid_contacts: list[ContactRecord]
    errors: list[ImportRowError]

    @property
    def contacts(self) -> list[ContactRecord]:
        """Backward-compatible alias for valid_contacts."""
        return self.valid_contacts


# ---------------------------------------------------------------------------
# Format Detectors & Sheet Readers
# ---------------------------------------------------------------------------

def inspect_spreadsheet(
    file_path: str | Path,
    sheet_name: str | None = None,
) -> SpreadsheetInspection:
    """
    Inspect a spreadsheet file, list sheets, detect headers, and propose mappings.
    Does not process rows yet.
    """
    path = Path(file_path)
    if not path.exists():
        raise SpreadsheetError(f"File not found: {path.resolve()}")

    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _inspect_csv(path)
    elif suffix == ".xlsx":
        return _inspect_xlsx(path, sheet_name)
    elif suffix == ".xls":
        return _inspect_xls(path, sheet_name)
    else:
        raise SpreadsheetError(
            f"Unsupported spreadsheet format '{suffix}'. Supported formats are: .csv, .xlsx, .xls"
        )


def _inspect_csv(path: Path) -> SpreadsheetInspection:
    try:
        raw_text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        try:
            raw_text = path.read_text(encoding="latin-1")
        except Exception as exc:
            raise SpreadsheetError(f"Could not read CSV file encoding: {exc}") from exc

    lines = [line for line in raw_text.splitlines() if line.strip()]
    if not lines:
        raise SpreadsheetError("CSV file is completely empty.")

    reader = csv.reader(lines)
    header_row_index = 1
    headers: list[str] = []

    for idx, row in enumerate(reader, start=1):
        clean_cells = [str(c).strip() for c in row if str(c).strip()]
        if clean_cells:
            headers = [str(c).strip() for c in row]
            header_row_index = idx
            break

    if not headers:
        raise SpreadsheetError("No header row found in CSV file.")

    mapping_res = auto_map_columns(headers)
    return SpreadsheetInspection(
        file_type="csv",
        sheet_names=["CSV"],
        selected_sheet="CSV",
        headers=headers,
        header_row_index=header_row_index,
        total_rows_approx=len(lines),
        proposed_mapping=mapping_res,
    )


def _inspect_xlsx(path: Path, sheet_name: str | None = None) -> SpreadsheetInspection:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise SpreadsheetError(f"Could not open XLSX workbook: {exc}") from exc

    sheet_names = wb.sheetnames
    if not sheet_names:
        wb.close()
        raise SpreadsheetError("XLSX workbook contains no sheets.")

    selected = sheet_name if sheet_name in sheet_names else sheet_names[0]
    ws = wb[selected]

    headers: list[str] = []
    header_row_index = 1
    total_rows = 0

    try:
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            total_rows += 1
            if not headers:
                clean_cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if clean_cells:
                    headers = [str(c).strip() if c is not None else f"Column_{i+1}" for i, c in enumerate(row)]
                    header_row_index = idx
    finally:
        wb.close()

    if not headers:
        raise SpreadsheetError(f"No non-empty rows found in sheet '{selected}'.")

    mapping_res = auto_map_columns(headers)
    return SpreadsheetInspection(
        file_type="xlsx",
        sheet_names=sheet_names,
        selected_sheet=selected,
        headers=headers,
        header_row_index=header_row_index,
        total_rows_approx=total_rows,
        proposed_mapping=mapping_res,
    )


def _inspect_xls(path: Path, sheet_name: str | None = None) -> SpreadsheetInspection:
    import xlrd

    try:
        wb = xlrd.open_workbook(str(path))
    except Exception as exc:
        raise SpreadsheetError(f"Could not open legacy XLS workbook: {exc}") from exc

    sheet_names = wb.sheet_names()
    if not sheet_names:
        raise SpreadsheetError("XLS workbook contains no sheets.")

    selected = sheet_name if sheet_name in sheet_names else sheet_names[0]
    ws = wb.sheet_by_name(selected)

    headers: list[str] = []
    header_row_index = 1
    total_rows = ws.nrows

    for r_idx in range(ws.nrows):
        row_vals = ws.row_values(r_idx)
        clean_cells = [str(c).strip() for c in row_vals if c is not None and str(c).strip()]
        if clean_cells:
            headers = [str(c).strip() if c is not None and str(c).strip() else f"Column_{i+1}" for i, c in enumerate(row_vals)]
            header_row_index = r_idx + 1
            break

    if not headers:
        raise SpreadsheetError(f"No non-empty rows found in sheet '{selected}'.")

    mapping_res = auto_map_columns(headers)
    return SpreadsheetInspection(
        file_type="xls",
        sheet_names=sheet_names,
        selected_sheet=selected,
        headers=headers,
        header_row_index=header_row_index,
        total_rows_approx=total_rows,
        proposed_mapping=mapping_res,
    )


# ---------------------------------------------------------------------------
# Complete Parsing with Mappings & Validation
# ---------------------------------------------------------------------------

def parse_spreadsheet_contacts(
    file_path: str | Path,
    mapping: dict[str, str | None] | None = None,
    sheet_name: str | None = None,
) -> SpreadsheetParseResult:
    """
    Parse a spreadsheet with given or auto-detected column mappings.
    Preserves 1-based source row numbers and original row order.
    Returns complete validation breakdown and ContactRecord objects.
    """
    path = Path(file_path)
    inspection = inspect_spreadsheet(path, sheet_name=sheet_name)

    # Resolve mapping
    if mapping:
        is_val, err, clean_map = validate_user_mapping(mapping, inspection.headers)
        if not is_val:
            raise SpreadsheetError(f"Invalid column mapping: {err}")
        active_map = clean_map
    else:
        # Use auto-detected mapping
        if not inspection.proposed_mapping.email_confident or not inspection.proposed_mapping.email_column:
            raise SpreadsheetError(
                "Could not automatically detect the Email column. Please specify column mapping."
            )
        active_map = {k: v for k, v in inspection.proposed_mapping.mapping.items() if v}

    raw_data_rows = _read_raw_rows(path, inspection)

    summary = SpreadsheetValidationSummary(
        total_source_rows=inspection.total_rows_approx,
        total_data_rows=len(raw_data_rows),
    )

    rows_result: list[RowValidationResult] = []
    valid_contacts: list[ContactRecord] = []
    errors: list[ImportRowError] = []

    seen_emails: dict[str, int] = {}  # normalized_email -> first_row_number
    duplicate_map: dict[str, DuplicateEmailRecord] = {}

    for row_dict, source_row_num in raw_data_rows:
        # Map values using active_map
        mapped_vals: dict[str, str] = {}
        for target_field in STANDARD_FIELDS:
            source_col = active_map.get(target_field)
            val = (row_dict.get(source_col) or "") if source_col else ""
            mapped_vals[target_field] = str(val).strip()

        raw_email = mapped_vals["email"]
        contact_name = mapped_vals["contact_name"]
        company_name = mapped_vals["company_name"]
        website = mapped_vals["website"]
        role_or_job_url = mapped_vals["role_or_job_url"]
        source_url = mapped_vals["source_url"]
        notes = mapped_vals["notes"]

        is_missing_company = not bool(company_name)
        is_missing_name = not bool(contact_name)

        if is_missing_company:
            summary.missing_company_count += 1
        if is_missing_name:
            summary.missing_name_count += 1

        # Check 1: Missing email
        if not raw_email:
            summary.missing_email_count += 1
            err_msg = "Email address is missing or blank."
            res = RowValidationResult(
                row_number=source_row_num,
                raw_values=row_dict,
                email="",
                contact_name=contact_name,
                company_name=company_name,
                website=website,
                role_or_job_url=role_or_job_url,
                source_url=source_url,
                notes=notes,
                is_valid=False,
                status="MISSING_EMAIL",
                reason=err_msg,
                is_missing_company=is_missing_company,
                is_missing_name=is_missing_name,
            )
            rows_result.append(res)
            errors.append(ImportRowError(row_number=source_row_num, raw_row=row_dict, reason=err_msg))
            continue

        # Check 2: Email syntax validation
        email_res = parse_email(raw_email)
        if not email_res.is_syntax_valid:
            summary.invalid_syntax_count += 1
            err_msg = f"Invalid email syntax: {email_res.parse_notes or 'Malformed format'}"
            res = RowValidationResult(
                row_number=source_row_num,
                raw_values=row_dict,
                email=raw_email,
                contact_name=contact_name,
                company_name=company_name,
                website=website,
                role_or_job_url=role_or_job_url,
                source_url=source_url,
                notes=notes,
                is_valid=False,
                status="SYNTAX_ERROR",
                reason=err_msg,
                is_missing_company=is_missing_company,
                is_missing_name=is_missing_name,
            )
            rows_result.append(res)
            errors.append(ImportRowError(row_number=source_row_num, raw_row=row_dict, reason=err_msg))
            continue

        norm_email = email_res.normalized_email
        is_free = (
            email_res.registered_domain in _FREE_PROVIDERS
            or email_res.email_domain in _FREE_PROVIDERS
        )
        if is_free:
            summary.free_provider_count += 1

        # Check 3: Duplicate email check
        if norm_email in seen_emails:
            first_row = seen_emails[norm_email]
            summary.duplicate_email_count += 1

            if norm_email not in duplicate_map:
                duplicate_map[norm_email] = DuplicateEmailRecord(
                    email=norm_email,
                    first_row=first_row,
                    duplicate_rows=[source_row_num],
                )
            else:
                duplicate_map[norm_email].duplicate_rows.append(source_row_num)

            err_msg = f"Duplicate email address (first seen at row {first_row}). Skipped."
            res = RowValidationResult(
                row_number=source_row_num,
                raw_values=row_dict,
                email=norm_email,
                contact_name=contact_name,
                company_name=company_name,
                website=website,
                role_or_job_url=role_or_job_url,
                source_url=source_url,
                notes=notes,
                is_valid=False,
                status="DUPLICATE",
                reason=err_msg,
                is_free_provider=is_free,
                is_missing_company=is_missing_company,
                is_missing_name=is_missing_name,
            )
            rows_result.append(res)
            errors.append(ImportRowError(row_number=source_row_num, raw_row=row_dict, reason=err_msg))
            continue

        # Record email seen
        seen_emails[norm_email] = source_row_num
        summary.valid_email_count += 1

        contact = ContactRecord(
            row_number=source_row_num,
            company_name=company_name,
            contact_name=contact_name,
            email_result=email_res,
            website=website,
            role_or_job_url=role_or_job_url,
            source_url=source_url,
            notes=notes,
        )

        res = RowValidationResult(
            row_number=source_row_num,
            raw_values=row_dict,
            email=norm_email,
            contact_name=contact_name,
            company_name=company_name,
            website=website,
            role_or_job_url=role_or_job_url,
            source_url=source_url,
            notes=notes,
            is_valid=True,
            status="VALID",
            reason="Syntax verified & de-duplicated",
            is_free_provider=is_free,
            is_missing_company=is_missing_company,
            is_missing_name=is_missing_name,
            contact_record=contact,
        )
        rows_result.append(res)
        valid_contacts.append(contact)

    summary.duplicate_details = list(duplicate_map.values())

    return SpreadsheetParseResult(
        inspection=inspection,
        applied_mapping=active_map,
        summary=summary,
        rows=rows_result,
        valid_contacts=valid_contacts,
        errors=errors,
    )


def _read_raw_rows(path: Path, inspection: SpreadsheetInspection) -> list[tuple[dict[str, str], int]]:
    """
    Read all rows from the spreadsheet after the header row.
    Returns list of (row_dict, 1_based_source_row_number).
    Skips completely blank rows.
    """
    suffix = inspection.file_type
    headers = inspection.headers
    header_idx = inspection.header_row_index
    out: list[tuple[dict[str, str], int]] = []

    if suffix == "csv":
        try:
            raw_text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            raw_text = path.read_text(encoding="latin-1")

        lines = raw_text.splitlines()
        reader = csv.reader(lines)
        for row_num, row_cells in enumerate(reader, start=1):
            if row_num <= header_idx:
                continue
            # Skip completely empty rows
            if not any(str(c).strip() for c in row_cells):
                continue
            row_dict: dict[str, str] = {}
            for i, h in enumerate(headers):
                row_dict[h] = str(row_cells[i]).strip() if i < len(row_cells) and row_cells[i] is not None else ""
            out.append((row_dict, row_num))

    elif suffix == "xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        try:
            ws = wb[inspection.selected_sheet]
            for row_num, row_cells in enumerate(ws.iter_rows(values_only=True), start=1):
                if row_num <= header_idx:
                    continue
                # Skip completely empty rows
                if not any(c is not None and str(c).strip() for c in row_cells):
                    continue
                row_dict = {}
                for i, h in enumerate(headers):
                    val = row_cells[i] if i < len(row_cells) else None
                    row_dict[h] = str(val).strip() if val is not None else ""
                out.append((row_dict, row_num))
        finally:
            wb.close()

    elif suffix == "xls":
        import xlrd

        wb = xlrd.open_workbook(str(path))
        ws = wb.sheet_by_name(inspection.selected_sheet)
        for r_idx in range(ws.nrows):
            row_num = r_idx + 1
            if row_num <= header_idx:
                continue
            row_vals = ws.row_values(r_idx)
            if not any(c is not None and str(c).strip() for c in row_vals):
                continue
            row_dict = {}
            for i, h in enumerate(headers):
                val = row_vals[i] if i < len(row_vals) else None
                # Handle numeric cells gracefully in xlrd
                if isinstance(val, float) and val.is_integer():
                    val_str = str(int(val))
                elif val is not None:
                    val_str = str(val).strip()
                else:
                    val_str = ""
                row_dict[h] = val_str
            out.append((row_dict, row_num))

    return out
