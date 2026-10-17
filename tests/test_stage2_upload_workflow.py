"""
tests/test_stage2_upload_workflow.py
------------------------------------
Comprehensive automated test suite for SmartMail Stage 2:
CV-First Upload Workflow, Spreadsheet Import, Column Auto-Mapping,
Pre-flight Contact Diagnostics, and Import Processing.
"""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Generator

import openpyxl
import pytest
import xlwt
from flask.testing import FlaskClient
from pypdf import PdfWriter

from src.contacts.column_mapper import (
    STANDARD_FIELDS,
    auto_map_columns,
    validate_user_mapping,
)
from src.contacts.spreadsheet_importer import (
    inspect_spreadsheet,
    parse_spreadsheet_contacts,
    SpreadsheetError,
)
from src.outreach.models import WorkflowMode
from src.outreach.send_log import SendLog
from src.ui.app import create_app
from src.ui.cv_validator import validate_cv_file, MAX_CV_SIZE_BYTES
from src.ui.storage import (
    get_registered_cv_metadata,
    register_uploaded_cv,
    store_and_register_cv,
    clear_registered_cv,
)
from src.workflow.engine import OutreachWorkflow


# ---------------------------------------------------------------------------
# Test Fixtures & Helpers
# ---------------------------------------------------------------------------

def create_valid_pdf_bytes(num_pages: int = 1) -> bytes:
    """Generate a minimal valid PDF in memory using pypdf."""
    writer = PdfWriter()
    for _ in range(num_pages):
        writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def create_xlsx_bytes(sheets_data: dict[str, list[list[str]]]) -> bytes:
    """Generate an .xlsx file with specified sheets and rows."""
    wb = openpyxl.Workbook()
    first = True
    for sheet_name, rows in sheets_data.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(title=sheet_name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_xls_bytes(sheets_data: dict[str, list[list[str]]]) -> bytes:
    """Generate a legacy .xls file with specified sheets and rows using xlwt."""
    wb = xlwt.Workbook(encoding="utf-8")
    for sheet_name, rows in sheets_data.items():
        ws = wb.add_sheet(sheet_name)
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                ws.write(r_idx, c_idx, val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def test_send_log(tmp_path: Path) -> SendLog:
    db_file = tmp_path / "stage2_send_log.db"
    return SendLog(db_path=db_file)


@pytest.fixture
def client(tmp_path: Path, test_send_log: SendLog) -> Generator[FlaskClient, None, None]:
    # Ensure reference CV exists for uploaded CV mode
    dummy_cv = tmp_path / "candidate_cv.pdf"
    dummy_cv.write_bytes(create_valid_pdf_bytes(1))
    register_uploaded_cv(dummy_cv)

    app = create_app({
        "TESTING": True,
        "SEND_LOG_PATH": str(test_send_log._db_path),
        "WTF_CSRF_ENABLED": False,
    })

    with app.test_client() as test_client:
        yield test_client


# ---------------------------------------------------------------------------
# 1. CV Upload & Validation Tests
# ---------------------------------------------------------------------------

class TestCVValidation:
    def test_valid_pdf(self, tmp_path: Path):
        pdf_path = tmp_path / "valid.pdf"
        pdf_path.write_bytes(create_valid_pdf_bytes(2))

        result = validate_cv_file(pdf_path)
        assert result.is_valid is True
        assert result.page_count == 2
        assert result.size_bytes > 0
        assert len(result.sha256) == 64
        assert result.error is None

    def test_empty_pdf(self, tmp_path: Path):
        pdf_path = tmp_path / "empty.pdf"
        pdf_path.write_bytes(b"")

        result = validate_cv_file(pdf_path)
        assert result.is_valid is False
        assert "empty" in result.error.lower()

    def test_corrupted_pdf(self, tmp_path: Path):
        pdf_path = tmp_path / "corrupted.pdf"
        pdf_path.write_bytes(b"%PDF-1.4 corrupt content that cannot be parsed by any reader")

        result = validate_cv_file(pdf_path)
        assert result.is_valid is False
        assert result.error is not None

    def test_non_pdf_file(self, tmp_path: Path):
        txt_path = tmp_path / "test.txt"
        txt_path.write_text("Hello world")

        result = validate_cv_file(txt_path)
        assert result.is_valid is False
        assert "pdf" in result.error.lower()

    def test_oversized_pdf(self, tmp_path: Path):
        large_path = tmp_path / "large.pdf"
        large_path.write_bytes(b"%PDF-1.4 " + (b"0" * (MAX_CV_SIZE_BYTES + 10)))

        result = validate_cv_file(large_path)
        assert result.is_valid is False
        assert "exceeds" in result.error.lower()

    def test_cv_storage_isolation_and_safe_filenames(self, tmp_path: Path):
        orig_bytes = create_valid_pdf_bytes(1)
        meta = store_and_register_cv(
            file_bytes=orig_bytes,
            original_filename="My Resume (Final) #1.pdf",
            output_dir=tmp_path / "uploads",
        )
        assert meta["registered"] is True
        assert meta["exists"] is True
        stored_path = Path(meta["path"])
        assert stored_path.exists()
        # Verify original bytes are intact (byte-for-byte fidelity)
        assert stored_path.read_bytes() == orig_bytes
        # Verify filename is sanitized and has unique prefix
        assert "My_Resume__Final___1" in stored_path.name
        assert stored_path.name.startswith("cv_")


# ---------------------------------------------------------------------------
# 2. Column Auto-Mapping Tests
# ---------------------------------------------------------------------------

class TestColumnMapping:
    def test_auto_map_standard_headers(self):
        headers = ["Email", "Name", "Company", "Website", "Role", "Source", "Notes"]
        res = auto_map_columns(headers)
        assert res.email_confident is True
        assert res.mapping["email"] == "Email"
        assert res.mapping["contact_name"] == "Name"
        assert res.mapping["company_name"] == "Company"
        assert res.mapping["website"] == "Website"
        assert res.mapping["role_or_job_url"] == "Role"
        assert res.mapping["source_url"] == "Source"
        assert res.mapping["notes"] == "Notes"

    def test_auto_map_aliases_and_case_variations(self):
        headers = ["  Contact Email  ", "RECIPIENT", "Organization", "Job Link", "Profile URL", "Comments"]
        res = auto_map_columns(headers)
        assert res.email_confident is True
        assert res.mapping["email"] == "Contact Email"
        assert res.mapping["contact_name"] == "RECIPIENT"
        assert res.mapping["company_name"] == "Organization"
        assert res.mapping["role_or_job_url"] == "Job Link"
        assert res.mapping["source_url"] == "Profile URL"
        assert res.mapping["notes"] == "Comments"

    def test_missing_email_header(self):
        headers = ["Name", "Company", "Phone", "Location"]
        res = auto_map_columns(headers)
        assert res.email_confident is False
        assert res.email_column is None
        assert res.mapping["contact_name"] == "Name"
        assert res.mapping["company_name"] == "Company"

    def test_validate_column_mapping_duplicate_rejection(self):
        headers = ["Email", "Name", "Company"]
        # Duplicate mapping: mapping two fields to the same column
        bad_mapping = {"email": "Email", "contact_name": "Email"}
        is_valid, err, clean = validate_user_mapping(bad_mapping, headers)
        assert is_valid is False
        assert "cannot be mapped to both" in err.lower()

    def test_validate_column_mapping_missing_email(self):
        headers = ["Name", "Company"]
        bad_mapping = {"contact_name": "Name", "company_name": "Company"}
        is_valid, err, clean = validate_user_mapping(bad_mapping, headers)
        assert is_valid is False
        assert "email column must be selected" in err.lower()

    def test_validate_column_mapping_nonexistent_column(self):
        headers = ["Email", "Name"]
        bad_mapping = {"email": "Email", "contact_name": "NonExistentColumn"}
        is_valid, err, clean = validate_user_mapping(bad_mapping, headers)
        assert is_valid is False
        assert "is not in the spreadsheet headers" in err.lower()


# ---------------------------------------------------------------------------
# 3. Spreadsheet Import & Row Order Preservation Tests
# ---------------------------------------------------------------------------

class TestSpreadsheetImport:
    def test_csv_import_row_numbers_and_order(self, tmp_path: Path):
        csv_path = tmp_path / "contacts.csv"
        rows = [
            ["Email", "Name", "Company"],
            ["alice@example.com", "Alice Smith", "Acme Inc"],
            ["bob@example.com", "Bob Jones", "Beta Corp"],
            ["charlie@example.com", "Charlie Brown", "Gamma LLC"],
        ]
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerows(rows)

        import_res = parse_spreadsheet_contacts(
            csv_path,
            mapping={"email": "Email", "contact_name": "Name", "company_name": "Company"},
        )

        assert len(import_res.valid_contacts) == 3
        # Check source row numbers (1-indexed: header is row 1, data rows are 2, 3, 4)
        assert import_res.valid_contacts[0].row_number == 2
        assert import_res.valid_contacts[0].email == "alice@example.com"
        assert import_res.valid_contacts[1].row_number == 3
        assert import_res.valid_contacts[1].email == "bob@example.com"
        assert import_res.valid_contacts[2].row_number == 4
        assert import_res.valid_contacts[2].email == "charlie@example.com"

    def test_xlsx_import_and_sheet_selection(self, tmp_path: Path):
        xlsx_path = tmp_path / "multisheet.xlsx"
        sheets = {
            "Candidates": [
                ["Contact Email", "Candidate Name"],
                ["dev1@startup.io", "Dev One"],
            ],
            "Partners": [
                ["Partner Email", "Partner Name", "Company"],
                ["partner@bigcorp.com", "Exec Partner", "Big Corp"],
            ],
        }
        xlsx_path.write_bytes(create_xlsx_bytes(sheets))

        # Check sheet listing via inspection
        inspection = inspect_spreadsheet(xlsx_path)
        assert inspection.sheet_names == ["Candidates", "Partners"]

        # Parse second sheet specifically
        import_res = parse_spreadsheet_contacts(
            xlsx_path,
            sheet_name="Partners",
            mapping={"email": "Partner Email", "contact_name": "Partner Name", "company_name": "Company"},
        )
        assert len(import_res.valid_contacts) == 1
        assert import_res.valid_contacts[0].email == "partner@bigcorp.com"
        assert import_res.valid_contacts[0].contact_name == "Exec Partner"
        assert import_res.valid_contacts[0].company_name == "Big Corp"
        assert import_res.valid_contacts[0].row_number == 2

    def test_legacy_xls_import(self, tmp_path: Path):
        xls_path = tmp_path / "legacy.xls"
        sheets = {
            "OldContacts": [
                ["Email", "Name", "Org"],
                ["oldie@vintage.org", "Oldie", "Vintage Org"],
            ]
        }
        xls_path.write_bytes(create_xls_bytes(sheets))

        inspection = inspect_spreadsheet(xls_path)
        assert "OldContacts" in inspection.sheet_names

        import_res = parse_spreadsheet_contacts(
            xls_path,
            sheet_name="OldContacts",
            mapping={"email": "Email", "contact_name": "Name", "company_name": "Org"},
        )
        assert len(import_res.valid_contacts) == 1
        assert import_res.valid_contacts[0].email == "oldie@vintage.org"
        assert import_res.valid_contacts[0].company_name == "Vintage Org"
        assert import_res.valid_contacts[0].row_number == 2

    def test_blank_rows_preserved_row_numbering(self, tmp_path: Path):
        csv_path = tmp_path / "blanks.csv"
        rows = [
            ["Email", "Name"],
            ["first@test.com", "First"],
            ["", ""],  # row 3: blank row
            ["third@test.com", "Third"],
        ]
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerows(rows)

        import_res = parse_spreadsheet_contacts(csv_path, mapping={"email": "Email", "contact_name": "Name"})
        # 2 non-blank valid contacts
        assert len(import_res.valid_contacts) == 2
        # Verify row 4 is preserved as row 4, not shifted to 3
        assert import_res.valid_contacts[0].row_number == 2
        assert import_res.valid_contacts[1].row_number == 4


# ---------------------------------------------------------------------------
# 4. Pre-Flight Contact Diagnostics & Validation Preview Tests
# ---------------------------------------------------------------------------

class TestValidationDiagnostics:
    def test_diagnostics_categories(self, tmp_path: Path):
        csv_path = tmp_path / "test_diag.csv"
        rows = [
            ["Email", "Name", "Company"],
            ["valid1@acme.com", "Alice", "Acme"],         # row 2: Valid
            ["not-an-email", "Bob", "Acme"],              # row 3: Invalid syntax
            ["", "Charlie", "NoEmail Inc"],               # row 4: Missing email
            ["valid1@acme.com", "Alice Duplicate", ""],   # row 5: Duplicate email + missing company
            ["lead@gmail.com", "", "Google Inc"],         # row 6: Free provider + missing name
        ]
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerows(rows)

        import_res = parse_spreadsheet_contacts(
            csv_path,
            mapping={"email": "Email", "contact_name": "Name", "company_name": "Company"},
        )
        metrics = import_res.summary

        # Check total data rows
        assert metrics.total_data_rows == 5

        # Check Mutually Exclusive Counts sum to total data rows:
        assert (
            metrics.valid_email_count
            + metrics.invalid_syntax_count
            + metrics.missing_email_count
            + metrics.duplicate_email_count
        ) == metrics.total_data_rows

        assert metrics.valid_email_count == 2  # valid1@acme.com (first occurrence), lead@gmail.com
        assert metrics.invalid_syntax_count == 1  # not-an-email
        assert metrics.missing_email_count == 1  # empty
        assert metrics.duplicate_email_count == 1  # second valid1@acme.com

        # Check Overlapping Informational Counts:
        assert metrics.missing_company_count == 1  # row 5 has empty company
        assert metrics.missing_name_count == 1  # row 6 has empty name
        assert metrics.free_provider_count == 1  # row 6 gmail.com

        # Check duplicate row attribution
        dup_row = [r for r in import_res.rows if r.row_number == 5][0]
        assert dup_row.status == "DUPLICATE"
        assert "first seen at row 2" in dup_row.reason

    def test_non_discarding_of_missing_company_or_free_provider(self, tmp_path: Path):
        """Missing company or free email provider must NOT discard contact."""
        csv_path = tmp_path / "valid_free.csv"
        rows = [
            ["Email", "Name", "Company"],
            ["freelancer@gmail.com", "John Doe", ""],
        ]
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerows(rows)

        import_res = parse_spreadsheet_contacts(
            csv_path,
            mapping={"email": "Email", "contact_name": "Name", "company_name": "Company"},
        )
        assert len(import_res.valid_contacts) == 1
        assert import_res.valid_contacts[0].email == "freelancer@gmail.com"
        assert import_res.valid_contacts[0].company_name == ""


# ---------------------------------------------------------------------------
# 5. UI Endpoints: Full 4-Step Wizard Workflow Tests
# ---------------------------------------------------------------------------

class TestUIWorkflowEndpoints:
    def test_cv_upload_ajax_and_clear(self, client: FlaskClient):
        valid_pdf = create_valid_pdf_bytes(1)

        # 1. Upload valid CV
        res = client.post(
            "/cv/upload",
            data={"cv_file": (io.BytesIO(valid_pdf), "my_cv.pdf")},
            headers={"X-Requested-With": "XMLHttpRequest"},
            content_type="multipart/form-data",
        )
        assert res.status_code == 200
        data = res.get_json()
        assert data["success"] is True
        assert data["cv"]["filename"] == "my_cv.pdf"
        assert data["cv"]["page_count"] == 1

        # Check registered CV state
        cv_meta = get_registered_cv_metadata()
        assert cv_meta["registered"] is True

        # 2. Clear CV
        res_clear = client.post("/cv/clear", headers={"X-Requested-With": "XMLHttpRequest"})
        assert res_clear.status_code == 200
        data_clear = res_clear.get_json()
        assert data_clear["success"] is True

        # Verify cleared
        assert get_registered_cv_metadata()["registered"] is False

    def test_cv_upload_invalid_file_rejected(self, client: FlaskClient):
        res = client.post(
            "/cv/upload",
            data={"cv_file": (io.BytesIO(b"Not a PDF"), "malformed.txt")},
            headers={"X-Requested-With": "XMLHttpRequest"},
            content_type="multipart/form-data",
        )
        assert res.status_code == 400
        data = res.get_json()
        assert data["success"] is False
        assert "error" in data

    def test_contacts_inspect_and_preview_workflow(self, client: FlaskClient):
        # Create CSV
        csv_bytes = b"Email,Name,Company\nalice@acme.com,Alice,Acme Corp\nbob@beta.org,Bob,Beta LLC\n"

        # 1. Inspect endpoint
        inspect_res = client.post(
            "/contacts/inspect",
            data={"contact_file": (io.BytesIO(csv_bytes), "leads.csv")},
            content_type="multipart/form-data",
        )
        assert inspect_res.status_code == 200
        inspect_data = inspect_res.get_json()
        assert inspect_data["success"] is True
        assert inspect_data["file_token"] is not None
        assert inspect_data["headers"] == ["Email", "Name", "Company"]
        assert inspect_data["proposed_mapping"]["email"] == "Email"

        file_token = inspect_data["file_token"]

        # 2. Preview endpoint
        preview_res = client.post(
            "/contacts/preview",
            json={
                "file_token": file_token,
                "sheet_name": "",
                "mapping": {
                    "email": "Email",
                    "contact_name": "Name",
                    "company_name": "Company",
                },
            },
        )
        assert preview_res.status_code == 200
        preview_data = preview_res.get_json()
        assert preview_data["success"] is True
        assert preview_data["summary"]["total_data_rows"] == 2
        assert preview_data["summary"]["valid_email_count"] == 2
        assert len(preview_data["rows"]) == 2

        # 3. Download validation report CSV
        report_res = client.post(
            "/contacts/download-report",
            json={
                "file_token": file_token,
                "sheet_name": "",
                "mapping": {
                    "email": "Email",
                    "contact_name": "Name",
                    "company_name": "Company",
                },
            },
        )
        assert report_res.status_code == 200
        assert report_res.content_type.startswith("text/csv")
        csv_text = report_res.data.decode("utf-8")
        assert "Source Row,Email,Name,Company,Website,Status,Explanation" in csv_text
        assert "alice@acme.com" in csv_text

    def test_import_process_dry_run_preserves_contacts(self, client: FlaskClient):
        # Setup CV first
        valid_pdf = create_valid_pdf_bytes(1)
        client.post(
            "/cv/upload",
            data={"cv_file": (io.BytesIO(valid_pdf), "resume.pdf")},
            content_type="multipart/form-data",
        )

        # Upload contact file to get file_token
        csv_content = b"Contact Email,Full Name,Organization\nlead1@site.com,Lead One,Site Inc\nlead2@site.com,Lead Two,Site Inc\n"
        inspect_res = client.post(
            "/contacts/inspect",
            data={"contacts_file": (io.BytesIO(csv_content), "site_leads.csv")},
            content_type="multipart/form-data",
        )
        file_token = inspect_res.get_json()["file_token"]

        # Process with custom mapping and dry_run=True
        mapping_dict = {
            "email": "Contact Email",
            "contact_name": "Full Name",
            "company_name": "Organization",
        }
        process_res = client.post(
            "/import/process",
            data={
                "file_token": file_token,
                "mapping": json.dumps(mapping_dict),
                "dry_run": "on",
            },
            follow_redirects=True,
        )
        assert process_res.status_code == 200
        page_html = process_res.data.decode("utf-8")
        # Check that redirect to dashboard or batch page occurred with success
        assert "Dashboard" in page_html or "Batch" in page_html or "Import" in page_html


# ---------------------------------------------------------------------------
# 6. Backend Regression Tests (OutreachWorkflow)
# ---------------------------------------------------------------------------

class TestWorkflowEngineRegression:
    def test_existing_process_csv_backward_compatibility(self, tmp_path: Path):
        """Ensure OutreachWorkflow.process_csv still functions with standard CSV files."""
        csv_path = tmp_path / "legacy_test.csv"
        rows = [
            ["email", "name", "company"],
            ["test1@example.com", "Test User 1", "Acme"],
            ["test2@example.com", "Test User 2", "Beta"],
        ]
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerows(rows)

        # Valid CV
        cv_path = tmp_path / "cv.pdf"
        cv_path.write_bytes(create_valid_pdf_bytes(1))

        send_log = SendLog(db_path=tmp_path / "regress_log.db")
        workflow = OutreachWorkflow(
            send_log=send_log,
            default_mode=WorkflowMode.UPLOADED_CV,
            uploaded_cv_path=cv_path,
        )

        batch = workflow.process_csv(csv_path, dry_run=True)
        assert batch is not None
        assert batch.total_rows == 2
        assert len(batch.results) == 2
        assert batch.results[0].contact_email == "test1@example.com"

